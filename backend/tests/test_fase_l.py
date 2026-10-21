"""Fase L: o motor serve sozinho os 4 pontos que eram do kernel-llm (ingestão de arquivo,
teste de modelo, teste de fonte de dados, extração de memórias) e o `graph` é desligável
por empresa, com backup e sem apagar nada."""

import asyncio
import gzip
import io
import json
import uuid

import httpx
import psycopg
import pytest
from app import embeddings, ingestion, memory_extraction, model_probe
from app.config import settings
from tests.conftest import auth


def _sql(query: str, params=()):
    with psycopg.connect(settings.database_url, autocommit=True) as conn:
        cur = conn.execute(query, params)
        return cur.fetchall() if cur.description else None


def _mock(handler):
    return httpx.AsyncClient(transport=httpx.MockTransport(handler))


# ------------------------------------------------------------- unitários ----


def test_chunk_text_overlap_and_end():
    text = "palavra " * 600
    chunks = ingestion.chunk_text(text)
    assert len(chunks) >= 4
    assert all(len(c) <= ingestion.CHUNK_CHARS for c in chunks)
    assert chunks[0][-ingestion.CHUNK_OVERLAP :] in chunks[1]
    assert ingestion.chunk_text("   \n ") == []
    assert ingestion.chunk_text("curto") == ["curto"]


def test_extract_text_formats():
    assert "olá" in ingestion.extract_text("olá, mundo".encode(), "text/plain", "a.txt")
    from docx import Document

    buf = io.BytesIO()
    doc = Document()
    doc.add_paragraph("Política de reembolso")
    doc.save(buf)
    assert "reembolso" in ingestion.extract_text(buf.getvalue(), "", "p.docx")
    from openpyxl import Workbook

    wb = Workbook()
    wb.active.append(["produto", 10])
    buf = io.BytesIO()
    wb.save(buf)
    assert "produto | 10" in ingestion.extract_text(buf.getvalue(), "", "t.xlsx")
    import fitz

    pdf = fitz.open()
    pdf.new_page().insert_text((72, 72), "Contrato XPTO")
    assert "Contrato XPTO" in ingestion.extract_text(pdf.tobytes(), "application/pdf", "c.pdf")


def test_embeddings_stub_is_deterministic_and_768():
    [a, b] = asyncio.run(embeddings.embed_texts({"provider": "stub"}, ["gato preto", "gato preto"]))
    assert a == b and len(a) == embeddings.DIMENSIONS


def test_embeddings_gemini_and_openai_request_shape():
    seen = []

    def handler(req: httpx.Request) -> httpx.Response:
        body = json.loads(req.content)
        seen.append((str(req.url), dict(req.headers), body))
        if "googleapis" in str(req.url):
            n = len(body["requests"])
            return httpx.Response(200, json={"embeddings": [{"values": [0.1] * 768}] * n})
        n = len(body["input"])
        data = [{"index": i, "embedding": [0.2] * 768} for i in reversed(range(n))]
        return httpx.Response(200, json={"data": data})

    async def go():
        async with _mock(handler) as c:
            g = await embeddings.embed_texts(
                {"provider": "gemini", "model": "gemini/gemini-embedding-001", "api_key": "K1"},
                ["a", "b"],
                client=c,
            )
            o = await embeddings.embed_texts(
                {"provider": "openai", "model": "text-embedding-3-small", "api_key": "K2"},
                ["a", "b", "c"],
                client=c,
            )
        return g, o

    g, o = asyncio.run(go())
    assert len(g) == 2 and len(o) == 3
    url, headers, body = seen[0]
    assert url.endswith("/models/gemini-embedding-001:batchEmbedContents")
    assert headers["x-goog-api-key"] == "K1"
    assert body["requests"][0]["outputDimensionality"] == 768
    assert seen[1][2]["dimensions"] == 768 and seen[1][1]["authorization"] == "Bearer K2"


def test_embeddings_provider_error_is_explicit():
    async def go():
        async with _mock(lambda r: httpx.Response(404, text="model not found")) as c:
            await embeddings.embed_texts({"provider": "openai", "api_key": "k"}, ["x"], client=c)

    with pytest.raises(embeddings.EmbeddingError, match="404"):
        asyncio.run(go())


@pytest.mark.parametrize(
    ("provider", "needle"),
    [
        ("gemini", "generateContent"),
        ("anthropic", "/messages"),
        ("openai", "api.openai.com/v1/chat/completions"),
        ("groq", "api.groq.com"),
        ("deepseek", "api.deepseek.com"),
        ("openai-compatible", "meu-host.test/v1/chat/completions"),
    ],
)
def test_probe_model_routes_each_provider(provider, needle):
    seen = []

    def handler(req: httpx.Request):
        seen.append(str(req.url))
        return httpx.Response(200, json={})

    spec = {"provider": provider, "model": "m", "api_key": "k", "api_base": None}
    if provider == "openai-compatible":
        spec["api_base"] = "https://meu-host.test/v1"
    ok, detail = asyncio.run(model_probe.probe_model(spec, client=_mock(handler)))
    assert (ok, detail) == (True, "")
    assert needle in seen[0]


def test_probe_model_failures_never_raise_and_hide_the_key():
    ok, detail = asyncio.run(
        model_probe.probe_model(
            {"provider": "openai", "model": "m", "api_key": "SEGREDO"},
            client=_mock(lambda r: httpx.Response(401, text="bad key SEGREDO")),
        )
    )
    assert ok is False and "401" in detail and "SEGREDO" not in detail

    def boom(req):
        raise httpx.ConnectError("sem rede")

    ok, detail = asyncio.run(
        model_probe.probe_model(
            {"provider": "gemini", "model": "m", "api_key": "k"}, client=_mock(boom)
        )
    )
    assert ok is False and "inacessível" in detail
    ok, detail = asyncio.run(model_probe.probe_model({"provider": "cobol", "model": "m"}))
    assert ok is False and "não suportado" in detail


# ----------------------------------------------------------- integração -----

integration = pytest.mark.integration


@pytest.fixture
def usuario(client, master_token, tenant):
    email = f"u-{uuid.uuid4().hex[:8]}@acme.com"
    r = client.post(
        "/api/users",
        json={
            "email": email,
            "name": "U",
            "password": "senha-forte-123",
            "tenant_id": tenant["id"],
        },
        headers=auth(master_token),
    )
    assert r.status_code == 201, r.text
    token = client.post(
        "/api/auth/login", json={"email": email, "password": "senha-forte-123"}
    ).json()["token"]
    return {"id": r.json()["id"], "token": token}


@integration
def test_upload_is_ingested_by_the_engine(client, tenant_admin):
    texto = b"Politica de reembolso da ACME: ate 30 dias. " * 100
    r = client.post(
        "/api/files",
        files={"file": ("politica.txt", io.BytesIO(texto), "text/plain")},
        headers=auth(tenant_admin["token"]),
    )
    assert r.status_code == 201, r.text
    fid = r.json()["id"]
    f = next(
        x
        for x in client.get("/api/files", headers=auth(tenant_admin["token"])).json()
        if x["id"] == fid
    )
    assert f["status"] == "ready" and f["chunk_count"] > 1 and f["error_detail"] is None
    n = _sql("SELECT count(*) FROM file_chunks WHERE file_id = %s", (fid,))[0][0]
    assert n == f["chunk_count"]
    # reprocessar substitui os trechos, sem duplicar
    assert (
        client.post(f"/api/files/{fid}/reprocess", headers=auth(tenant_admin["token"])).status_code
        == 200
    )
    assert _sql("SELECT count(*) FROM file_chunks WHERE file_id = %s", (fid,))[0][0] == n


@integration
def test_upload_without_text_ends_in_error_status(client, tenant_admin):
    r = client.post(
        "/api/files",
        files={"file": ("vazio.txt", io.BytesIO(b"   \n  "), "text/plain")},
        headers=auth(tenant_admin["token"]),
    )
    fid = r.json()["id"]
    row = _sql("SELECT status, error_detail FROM files WHERE id = %s", (fid,))[0]
    assert row[0] == "error" and "sem texto" in row[1]


@integration
def test_datasource_test_goes_through_the_harness(client, tenant_admin, fake_harness):
    h = auth(tenant_admin["token"])
    ok = client.post(
        "/api/datasources",
        json={"name": "vendas", "kind": "sqlite", "config": {"path": "/tmp/x.db"}},
        headers=h,
    )
    assert ok.status_code == 201, ok.text
    r = client.post(f"/api/datasources/{ok.json()['id']}/test", headers=h)
    assert r.json() == {"ok": True, "detail": ""}
    sent = fake_harness.calls["ds_tests"][0]
    assert sent["kind"] == "sqlite" and sent["config"] == {"path": "/tmp/x.db"}

    bad = client.post(
        "/api/datasources",
        json={"name": "ruim", "kind": "sqlite", "config": {"falha": True}},
        headers=h,
    ).json()
    r = client.post(f"/api/datasources/{bad['id']}/test", headers=h)
    assert r.json() == {"ok": False, "detail": "conexão recusada"}
    last = _sql("SELECT last_test_ok FROM datasources WHERE id = %s", (bad["id"],))[0][0]
    assert last is False


@integration
def test_datasource_test_with_harness_down_reports_instead_of_500(
    client, tenant_admin, monkeypatch
):
    monkeypatch.setattr(settings, "kernel_harness_url", "")
    h = auth(tenant_admin["token"])
    ds = client.post(
        "/api/datasources",
        json={"name": "x", "kind": "sqlite", "config": {"path": "/tmp/y.db"}},
        headers=h,
    ).json()
    r = client.post(f"/api/datasources/{ds['id']}/test", headers=h)
    assert r.status_code == 200 and r.json()["ok"] is False


@integration
def test_ai_service_test_calls_the_provider_not_the_kernel(client, tenant_admin, monkeypatch):
    chamadas = []

    async def fake_probe(spec, **_):
        chamadas.append(spec)
        return (spec["api_key"] == "chave-boa", "" if spec["api_key"] == "chave-boa" else "401")

    monkeypatch.setattr("app.routes.ai_services.probe_model", fake_probe)
    h = auth(tenant_admin["token"])
    svc = client.post(
        "/api/ai-services",
        json={
            "name": "g",
            "provider": "gemini",
            "model": "gemini-2.5-flash",
            "api_key": "chave-boa",
        },
        headers=h,
    ).json()
    assert client.post(f"/api/ai-services/{svc['id']}/test", headers=h).json() == {
        "ok": True,
        "detail": "",
    }
    assert chamadas[0]["provider"] == "gemini" and chamadas[0]["api_key"] == "chave-boa"
    client.put(f"/api/ai-services/{svc['id']}", json={"api_key": "chave-ruim"}, headers=h)
    r = client.post(f"/api/ai-services/{svc['id']}/test", headers=h).json()
    assert r["ok"] is False


def _chat_com_mensagens(tenant_id, user_id, kernel, pares):
    cid = str(uuid.uuid4())
    _sql(
        "INSERT INTO chats (id, tenant_id, user_id, title, kernel) VALUES (%s,%s,%s,%s,%s)",
        (cid, tenant_id, user_id, "t", kernel),
    )
    for i, (role, txt) in enumerate(pares):
        _sql(
            """INSERT INTO chat_messages (chat_id, role, content, created_at)
               VALUES (%s,%s,%s, now() + (%s || ' seconds')::interval)""",
            (cid, role, txt, i),
        )
    return cid


@integration
def test_memory_extraction_reads_messages_dedupes_and_keeps_watermark(
    client, tenant, usuario, fake_harness
):
    cid = _chat_com_mensagens(
        tenant["id"],
        usuario["id"],
        "harness",
        [("user", "eu trabalho na Acme"), ("assistant", "ok")],
    )
    fake_harness.calls["facts_to_return"] = ["Trabalha na Acme", "Prefere respostas curtas"]

    def rodar():
        return asyncio.run(
            memory_extraction.extract_memories(
                chat_id=cid,
                tenant_id=tenant["id"],
                user_id=usuario["id"],
                embedding={"provider": "stub"},
                model="agente-padrao",
            )
        )

    assert rodar() == 2
    enviado = fake_harness.calls["mem_extract"][0]
    assert [m["role"] for m in enviado["messages"]] == ["user", "assistant"]
    assert enviado["model"] == {"alias": "agente-padrao"}
    # sem mensagens novas: não chama o harness de novo (marca d'água)
    assert rodar() == 0 and len(fake_harness.calls["mem_extract"]) == 1
    # mensagem nova com o mesmo fato: deduplicado pelo embedding
    _sql("INSERT INTO chat_messages (chat_id, role, content) VALUES (%s,'user','de novo')", (cid,))
    assert rodar() == 0
    assert (
        _sql("SELECT messages_read FROM memory_extraction_state WHERE thread_id = %s", (cid,))[0][0]
        == 3
    )
    assert _sql("SELECT count(*) FROM memories WHERE tenant_id = %s", (tenant["id"],))[0][0] == 2


@integration
def test_memory_extraction_never_raises_when_harness_is_down(client, tenant, usuario, monkeypatch):
    monkeypatch.setattr(settings, "kernel_harness_url", "http://127.0.0.1:1")
    cid = _chat_com_mensagens(tenant["id"], usuario["id"], "harness", [("user", "oi")])
    saved = asyncio.run(
        memory_extraction.extract_memories(
            chat_id=cid,
            tenant_id=tenant["id"],
            user_id=usuario["id"],
            embedding={"provider": "stub"},
        )
    )
    assert saved == 0
    # não marcou como lida: tenta de novo no próximo turno
    assert not _sql("SELECT 1 FROM memory_extraction_state WHERE thread_id = %s", (cid,))


# ------------------------------------------------------------ L-2: graph ----


def _desligar_graph(client, master_token, tenant, on=False):
    r = client.put(
        f"/api/tenants/{tenant['id']}", json={"graph_enabled": on}, headers=auth(master_token)
    )
    assert r.status_code == 200, r.text
    return r.json()


@integration
def test_graph_flag_default_on_and_visible_to_master_and_tenant(
    client, master_token, tenant, tenant_admin
):
    items = client.get("/api/tenants", headers=auth(master_token)).json()["items"]
    assert next(t for t in items if t["id"] == tenant["id"])["graph_enabled"] is True
    assert (
        client.get("/api/auth/me", headers=auth(tenant_admin["token"])).json()[
            "graph_agents_enabled"
        ]
        is True
    )
    _desligar_graph(client, master_token, tenant)
    assert (
        client.get("/api/auth/me", headers=auth(tenant_admin["token"])).json()[
            "graph_agents_enabled"
        ]
        is False
    )


@integration
def test_graph_off_blocks_new_runs_and_templates_but_history_stays_readable(
    client, master_token, tenant, tenant_admin, fake_kernel, monkeypatch
):
    h = auth(tenant_admin["token"])
    monkeypatch.setattr(settings, "kernel_url", fake_kernel)
    # antes: template structured criável e conversa funciona
    t = client.post("/api/templates", json={"name": "legado"}, headers=h)
    assert t.status_code == 201, t.text
    with client.stream("POST", "/api/chat/send", json={"message": "oi"}, headers=h) as resp:
        body = "".join(resp.iter_text())
    assert "event: done" in body
    chat_id = client.get("/api/chats", headers=h).json()[0]["id"]

    _desligar_graph(client, master_token, tenant)

    # novas execuções: 403 (conversa nova e conversa antiga)
    for payload in ({"message": "outra"}, {"message": "continua", "chat_id": chat_id}):
        with client.stream("POST", "/api/chat/send", json=payload, headers=h) as resp:
            resp.read()
        assert resp.status_code == 403, resp.text
        assert "descontinuados" in resp.json()["detail"]
    # novo template / nova versão structured: 403
    assert client.post("/api/templates", json={"name": "novo"}, headers=h).status_code == 403
    assert (
        client.post(
            f"/api/templates/{t.json()['id']}/versions",
            json={"supervisor_prompt": "x", "agents": []},
            headers=h,
        ).status_code
        == 403
    )
    # a lista esconde os structured; o histórico segue legível
    assert all(
        x["agent_kind"] != "structured" for x in client.get("/api/templates", headers=h).json()
    )
    msgs = client.get(f"/api/chats/{chat_id}/messages", headers=h)
    assert msgs.status_code == 200 and len(msgs.json()) >= 2
    # reativar devolve tudo
    _desligar_graph(client, master_token, tenant, on=True)
    assert client.post("/api/templates", json={"name": "volta"}, headers=h).status_code == 201


@integration
def test_graph_off_does_not_touch_skills_agents(
    client, master_token, tenant, tenant_admin, fake_harness
):
    client.put(
        f"/api/tenants/{tenant['id']}", json={"harness_enabled": True}, headers=auth(master_token)
    )
    _desligar_graph(client, master_token, tenant)
    h = auth(tenant_admin["token"])
    r = client.post("/api/templates", json={"name": "sk", "agent_kind": "skills"}, headers=h)
    assert r.status_code == 201, r.text
    assert any(x["agent_kind"] == "skills" for x in client.get("/api/templates", headers=h).json())


@integration
def test_graph_shutdown_backs_up_quarantines_and_deletes_nothing(
    client, master_token, tenant, tenant_admin, usuario, tmp_path
):
    from app import graph_shutdown

    h = auth(tenant_admin["token"])
    t = client.post("/api/templates", json={"name": "legado"}, headers=h).json()
    cid = _chat_com_mensagens(
        tenant["id"], usuario["id"], "llm", [("user", "oi"), ("assistant", "olá")]
    )
    outro = _chat_com_mensagens(tenant["id"], usuario["id"], "harness", [("user", "skills")])
    _sql(
        "INSERT INTO memories (tenant_id, user_id, content) VALUES (%s,%s,'fato')",
        (tenant["id"], usuario["id"]),
    )
    _sql(
        """INSERT INTO tool_calls (tenant_id, chat_id, agent_name, tool_name, status)
           VALUES (%s,%s,'a','t','ok')""",
        (tenant["id"], cid),
    )

    sim = graph_shutdown.desligar(tenant["tenant_key"], tmp_path, simular=True)
    assert sim["simulado"] is True
    assert sim["contagens"]["chats"] == 1 and sim["contagens"]["chat_messages"] == 2
    assert sim["contagens"]["templates"] >= 1 and sim["contagens"]["memories"] == 1
    assert sim["contagens"]["tool_calls"] == 1
    flag = _sql("SELECT graph_enabled FROM tenants WHERE id = %s", (tenant["id"],))[0][0]
    assert flag is True  # simulação não liga a quarentena

    antes = {
        t_: _sql(f"SELECT count(*) FROM {t_}")[0][0]  # noqa: S608
        for t_ in ("chats", "chat_messages", "templates", "memories", "tool_calls")
    }
    out = graph_shutdown.desligar(tenant["tenant_key"], tmp_path)
    assert out["simulado"] is False
    assert _sql("SELECT graph_enabled FROM tenants WHERE id = %s", (tenant["id"],))[0][0] is False
    depois = {
        t_: _sql(f"SELECT count(*) FROM {t_}")[0][0]  # noqa: S608
        for t_ in antes
    }
    assert depois == antes  # nada apagado
    log = _sql(
        "SELECT action, backup_path, counts FROM graph_shutdown_log WHERE tenant_id = %s",
        (tenant["id"],),
    )
    assert log[-1][0] == "desligar" and log[-1][2]["chat_messages"] == 2

    # o backup é relido e o conteúdo bate
    from pathlib import Path

    pasta = Path(out["pasta"])
    with gzip.open(pasta / "chat_messages.jsonl.gz", "rt", encoding="utf-8") as fh:
        linhas = [json.loads(x) for x in fh]
    assert {x["content"] for x in linhas} == {"oi", "olá"}
    assert all(x["chat_id"] == cid for x in linhas) and outro not in {x["chat_id"] for x in linhas}
    manifesto = json.loads((pasta / "manifest.json").read_text(encoding="utf-8"))
    assert manifesto["arquivos"]["chats"]["linhas"] == 1
    graph_shutdown.verificar_backup(out)

    # backup adulterado é detectado
    (pasta / "chats.jsonl.gz").write_bytes(b"lixo")
    with pytest.raises(RuntimeError, match="corrompido"):
        graph_shutdown.verificar_backup(out)

    # a UI/API do dono continua lendo o histórico
    assert (
        client.get(f"/api/chats/{cid}/messages", headers=auth(usuario["token"])).status_code == 200
    )

    graph_shutdown.reativar(tenant["tenant_key"])
    assert _sql("SELECT graph_enabled FROM tenants WHERE id = %s", (tenant["id"],))[0][0] is True
    assert (
        _sql(
            "SELECT action FROM graph_shutdown_log WHERE tenant_id = %s "
            "ORDER BY created_at DESC LIMIT 1",
            (tenant["id"],),
        )[0][0]
        == "reativar"
    )
    assert t["id"]


@integration
def test_graph_shutdown_refuses_when_an_inbox_still_points_to_graph(
    client, tenant, tenant_admin, tmp_path
):
    from app import graph_shutdown

    h = auth(tenant_admin["token"])
    t = client.post("/api/templates", json={"name": "da-caixa"}, headers=h).json()
    _sql(
        "INSERT INTO inbox_ai_links (tenant_id, chatwoot_inbox_id, template_id) "
        "VALUES (%s, 77, %s)",
        (tenant["id"], t["id"]),
    )
    with pytest.raises(graph_shutdown.Bloqueio):
        graph_shutdown.desligar(tenant["tenant_key"], tmp_path)
    assert _sql("SELECT graph_enabled FROM tenants WHERE id = %s", (tenant["id"],))[0][0] is True
    assert not list(tmp_path.iterdir())  # recusou antes de gerar qualquer arquivo
    with pytest.raises(LookupError):
        graph_shutdown.desligar("nao-existe", tmp_path)
